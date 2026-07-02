import argparse
import csv
import json
import pathlib

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


OBJECT_TITLE = (
    "ГОРНО-ОБОГАТИТЕЛЬНЫЙ КОМПЛЕКС ПО ДОБЫЧЕ\n"
    "И ПЕРЕРАБОТКЕ ФЛЮОРИТ-БЕРИЛЛИЕВЫХ РУД\n"
    "МЕСТОРОЖДЕНИЯ «ЕРМАКОВСКОЕ»"
)

SYSTEM_TITLE = "Комплексная система безопасности\n(Система охранная телевизионная)"

CHECK_ROWS = [
    ("1", "Проверить соответствие установленного оборудования проекту", "Визуальная сверка с проектной и исполнительной документацией", "Оборудование соответствует проекту"),
    ("2", "Проверить маркировку", "Визуальная проверка маркировки кабелей, клемм, портов, устройств", "Маркировка соответствует документации"),
    ("3", "Проверить качество монтажа", "Визуальный осмотр креплений, корпусов, соединений", "Монтаж выполнен качественно, повреждений нет"),
    ("4", "Проверить электропитание", "Измерение / контроль наличия питающего напряжения", "Напряжение соответствует норме"),
    ("5", "Проверить линии связи и подключения", "Проверка правильности подключения цепей питания, управления и передачи данных", "Подключение выполнено правильно"),
    ("6", "Проверить первичное включение", "Подача питания в штатном режиме", "Устройство включается, аварийная индикация отсутствует"),
    ("7", "Проверить обмен с системой", "Контроль отображения устройства в системе / ПО", "Устройство определяется и доступно"),
    ("8", "Проверить корректность конфигурации", "Сверка параметров с проектом и принятой конфигурацией", "Параметры соответствуют проекту"),
    ("9", "Проверить регистрацию событий", "Инициирование штатного события", "Событие регистрируется в системе"),
    ("10", "Проверить работу после перезапуска питания", "Отключение и повторная подача питания", "Работоспособность восстанавливается"),
    ("11", "Проверить устойчивость работы", "Наблюдение в течение контрольного времени", "Сбоев и ложных срабатываний нет"),
    ("12", "Проверить журналы / сообщения системы", "Просмотр сообщений и событий в ПО", "Критические ошибки отсутствуют"),
    ("13", "Проверить наличие видеосигнала", "Просмотр изображения в системе видеонаблюдения", "Изображение устойчивое, доступно"),
    ("14", "Проверить качество изображения", "Контроль резкости, экспозиции, цветопередачи, зоны обзора", "Изображение соответствует условиям эксплуатации"),
    ("15", "Проверить запись и архивирование", "Контроль записи в архив / воспроизведение", "Запись выполняется корректно"),
    ("16", "Проверить реакцию на изменение освещенности", "Проверка в режимах день/ночь, при наличии", "Переключение режимов корректное"),
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, *, bold: bool = False, size: int = 9, align=None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text or "")
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{name}"))
        if element is None:
            element = OxmlElement(f"w:{name}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")


def set_column_widths(table, widths_cm):
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)


def compact_mac(value: str) -> str:
    value = (value or "").strip()
    if ":" in value:
        return value.upper()
    if len(value) == 12:
        return ":".join(value[i : i + 2] for i in range(0, 12, 2)).upper()
    return value


def json_value(row: dict, *names: str) -> str:
    for column in ("device_info_json", "network_json"):
        raw = row.get(column) or ""
        if not raw:
            continue
        try:
            data = json.loads(raw)
        except json.JSONDecodeError:
            continue
        found = find_first(data, {name.lower() for name in names})
        if found:
            return found
    return ""


def find_first(data, names: set[str]) -> str:
    if isinstance(data, dict):
        for key, value in data.items():
            if key.lower() in names and value not in ("", None, "null"):
                return str(value)
        for value in data.values():
            found = find_first(value, names)
            if found:
                return found
    elif isinstance(data, list):
        for item in data:
            found = find_first(item, names)
            if found:
                return found
    return ""


def load_camera_rows(paths: list[pathlib.Path], start_ip: str, end_ip: str) -> list[dict]:
    by_ip: dict[str, dict] = {}
    for path in paths:
        if not path.exists():
            continue
        with path.open("r", newline="", encoding="utf-8-sig") as file:
            for row in csv.DictReader(file):
                ip = (row.get("ip") or row.get("assigned_ip") or "").strip()
                if not ip:
                    continue
                current = by_ip.setdefault(ip, {"ip": ip})
                model = json_value(row, "DeviceModel", "productModel") or row.get("model") or ""
                serial = json_value(row, "SerialNumber", "serialNumber", "sn") or row.get("serial_number") or ""
                mac = json_value(row, "MAC", "MACAddress", "macAddress") or row.get("mac") or ""
                for key, value in (("model", model), ("serial_number", serial), ("mac", compact_mac(mac))):
                    if value and not current.get(key):
                        current[key] = value

    start = tuple(int(part) for part in start_ip.split("."))
    end = tuple(int(part) for part in end_ip.split("."))
    rows = []
    for ip, data in by_ip.items():
        parts = tuple(int(part) for part in ip.split("."))
        if start <= parts <= end:
            rows.append(data)
    return sorted(rows, key=lambda item: tuple(int(part) for part in item["ip"].split(".")))


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10)
    return doc


def add_centered_paragraph(doc: Document, text: str, *, size: int = 11, bold: bool = False, after_pt: int = 6) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(after_pt)
    for index, line in enumerate(text.splitlines()):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        run.bold = bold
        run.font.name = "Times New Roman"
        run.font.size = Pt(size)


def add_camera_table(doc: Document, rows: list[dict]) -> None:
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run("Сводная таблица видеокамер")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    table = doc.add_table(rows=1, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    headers = ["№", "IP-адрес", "Модель", "Серийный номер", "MAC", "Примечание"]
    widths = [0.9, 2.4, 4.9, 4.0, 3.0, 2.3]
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_text(cell, header, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(cell, "D9EAF7")
    set_repeat_table_header(table.rows[0])

    for index, camera in enumerate(rows, start=1):
        cells = table.add_row().cells
        values = [
            str(index),
            camera.get("ip", ""),
            camera.get("model", ""),
            camera.get("serial_number", ""),
            camera.get("mac", ""),
            "",
        ]
        for cell, value in zip(cells, values):
            align = WD_ALIGN_PARAGRAPH.CENTER if value in (str(index), camera.get("ip", ""), camera.get("mac", "")) else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cell, value, size=8, align=align)
    set_column_widths(table, widths)
    set_table_borders(table)


def add_equipment_fields(doc: Document, camera: dict) -> None:
    fields = [
        ("1. Наименование оборудования", "Видеокамера"),
        ("2. Тип / условное обозначение", camera.get("model", "")),
        ("3. Серийный номер", camera.get("serial_number", "")),
        ("4. Месторасположение", camera.get("location", "")),
        ("5. IP-адрес / адрес в системе / номер линии", camera.get("ip", "")),
    ]
    for label, value in fields:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(1)
        run = paragraph.add_run(f"{label}: ")
        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(10)
        value_run = paragraph.add_run(value or "______________________________")
        value_run.font.name = "Times New Roman"
        value_run.font.size = Pt(10)


def add_checklist(doc: Document, camera: dict, index: int) -> None:
    doc.add_page_break()
    add_centered_paragraph(
        doc,
        f"Чек-лист проверки конфигурации и системы № {index}",
        size=13,
        bold=True,
        after_pt=10,
    )
    add_equipment_fields(doc, camera)
    doc.add_paragraph()

    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    headers = ["№", "Проверка", "Действие", "Критерий", "Отметка"]
    widths = [0.8, 4.4, 5.5, 4.8, 2.0]
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_text(cell, header, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(cell, "D9EAF7")
    set_repeat_table_header(table.rows[0])

    for number, check, action, criterion in CHECK_ROWS:
        cells = table.add_row().cells
        values = [number, check, action, criterion, "☐ Да / ☐ Нет"]
        for cell, value in zip(cells, values):
            align = WD_ALIGN_PARAGRAPH.CENTER if cell in (cells[0], cells[4]) else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cell, value, size=8, align=align)
    set_column_widths(table, widths)
    set_table_borders(table)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("Примечание: ").bold = True
    p.add_run("состав проверок уточняется в зависимости от проектных решений и требований эксплуатационной документации изготовителя.")

    doc.add_paragraph()
    sign = doc.add_table(rows=2, cols=3)
    sign.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, value in zip(sign.rows[0].cells, ["Должность", "ФИО", "Подпись"]):
        set_cell_text(cell, value, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(cell, "D9EAF7")
    for cell in sign.rows[1].cells:
        set_cell_text(cell, "", size=9)
    set_column_widths(sign, [5.0, 7.0, 5.0])
    set_table_borders(sign)


def parse_args() -> argparse.Namespace:
    script_dir = pathlib.Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description="Build DOCX PNR camera report without images.")
    parser.add_argument("--start-ip", default="10.53.240.30")
    parser.add_argument("--end-ip", default="10.53.240.132")
    parser.add_argument("--checklist-count", type=int, default=20)
    parser.add_argument(
        "--include-without-serial",
        action="store_true",
        help="Include cameras without serial numbers. By default only cameras with serial numbers are used.",
    )
    parser.add_argument("--output-docx", default=str(script_dir / "camera_pnr_report.docx"))
    parser.add_argument(
        "--input-csv",
        action="append",
        default=[
            str(script_dir / "camera_info_after.csv"),
            str(script_dir / "camera_info_clean.csv"),
            str(script_dir / "camera_inventory.csv"),
        ],
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    rows = load_camera_rows([pathlib.Path(path) for path in args.input_csv], args.start_ip, args.end_ip)
    if not args.include_without_serial:
        rows = [row for row in rows if row.get("serial_number")]
    doc = setup_document()
    add_centered_paragraph(doc, OBJECT_TITLE, size=11, bold=True, after_pt=14)
    add_centered_paragraph(doc, SYSTEM_TITLE, size=11, after_pt=14)
    add_camera_table(doc, rows)
    for index, camera in enumerate(rows[: args.checklist_count], start=1):
        add_checklist(doc, camera, index)

    output_path = pathlib.Path(args.output_docx)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    print(f"Written {len(rows)} camera rows and {min(len(rows), args.checklist_count)} checklists to {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
